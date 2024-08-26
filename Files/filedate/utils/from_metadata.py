#!/usr/bin/env python

import re
#from pathlib import Path

from datetime import datetime
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from hachoir.parser import createParser
from hachoir.core import config as HachoirConfig
HachoirConfig.quiet = True
from hachoir.metadata import extractMetadata # creation_date as datetime
# also from EXIF (OffsetTime-EXIF v.2.31 - I don't know if it's taken into account)
# Many formats, but is there creation/modif. date-time?
# MP3, WAV, OGG, MIDI, AIFF, AIFC, RA
# GIF, JPEG, PCX, PNG, TGA, TIFF, WMF, XCF
# EXE	 # WMV, AVI, MKV, MOV, RM

import sys; sys.path.insert(1, sys.path[0]+'/..') # if `filedate` not installed
from filedate import FileDate

from dateutil.parser import parse #used if FromMetadata.VERBOSE

#-=-=-=-#

class FromMetadata:
	"""Utility for setting file date based on its metadata.
Poorly tested, but can be useful, 
for example, for saved attachments from e-mail.
"""
	VERBOSE = True
	LIST_ONLY = False # List only (in combination with VERBOSE) - don’t modify timestamp

	def __init__(self, file_path: str):
		self.file_path = file_path
		self._metadata = {} # as timestamp or str: {'created': dt_tm0, 'modified': dt_tm1, ...}

#-=-=-=-#

# MIT - based on https://github.com/JMousqueton/MetaLookup/ ver "0.3.1"
# pip install PyPDF2 python-docx openpyxl python-pptx hachoir

	def extract_metadata(self, add_modified_if_missing = True):
		"""add_modified_if_missing - and if created exists
		Mainly 'modified' will be used to correct the timestamp of the file"""

		def extract_pdf_metadata(pdf_path): #.pdf
			with open(pdf_path, 'rb') as file:
				pdf = PdfReader(file)
				try:
					info = pdf.trailer["/Info"]
					metadata = {key[1:]: info[key] for key in info if key != '/ID'}
					# CreationDate: D:20240528193019+00'00' # ModDate: D:20240528193019+00'00'
					# CreationDate: D:20210316065455Z00'00'
					metadata_ = {}
					for (k0, k1) in (('CreationDate', 'created'), ('ModDate', 'modified')):
						k0_ = metadata.get(k0)
						if k0_:
							metadata_[k1] = re.sub(r"Z.+$",r"Z",re.sub(r"'.+$", "", re.sub(r"^.+:","",k0_)))
							# created: 20240528193019+00, modified: 20240528193019+00
					for (k, v) in metadata_.items():
						v_ = None
						try:
							v_ = parse(v) #$#; print(f'{v_=}') #$#
							if v_.year < 1970: # eq. CreationDate: D:16010101000000Z
								v_ = None
						except:
							v_d = re.sub(r'(^.*?(\d{8,}).*$)',r'\2',v)[:14] # 20210316065455
							if FromMetadata.VERBOSE: print(f'    {v} -> {v_d}')
							try: 
								v_ = parse(v_d)
							except Exception as error:
								if FromMetadata.VERBOSE: print(f'{error}')
								v_ = None
						metadata_[k] = v_
					#
					metadata.update(metadata_)
					return add_missing_modified(metadata)
				except KeyError:
					return {}
		#
		def extract_office_metadata(doc_path): # .docx, .xlsx, .pptx
			if doc_path.endswith('.docx'):
				doc = Document(doc_path)
				core_props = doc.core_properties
			elif doc_path.endswith('.xlsx'):
				wb = load_workbook(doc_path)
				core_props = wb.properties
			elif doc_path.endswith('.pptx'):
				pres = Presentation(doc_path)
				core_props = pres.core_properties
			else:
				return {}
			#
			metadata = { # timezone.utc +00, eq. created: 2021-04-07 06:21:00+00:00 modified: 2023-10-03 05:14:00+00:00
				'created': getattr(core_props, 'created', None),
				'modified': getattr(core_props, 'modified', None),
			} 
			for k,v in metadata.items():
				if k in ('created', 'modified'):
				# m.in. In .xlsx and .pptx you can find dates without tzinfo, but they actually mean UTC +00
					if isinstance(v, datetime) and v.tzinfo is None: # or? v.tzinfo.utcoffset(v) is None:
						if FromMetadata.VERBOSE: print('    ### missing tzinfo +00 ###')
						#v.replace(tzinfo=timezone.utc) # probably doesn't work without: import pytz
						v = f'{v}' + '+00:00' 
						metadata[k] = v
			return add_missing_modified(metadata)
		#
		def extract_multimedia_metadata(video_path): # and all other
			parser = createParser(video_path)
			if not parser:
				if FromMetadata.VERBOSE: print(f'    Unable to parse video file: "{video_path}"')
				return {}
			with parser:
				try:
					metadata_ = extractMetadata(parser)
				except Exception as e:
					if FromMetadata.VERBOSE: print(f'    Metadata extraction error: "{e}"')
					return {}
			if not metadata_:
				return {}
			#
			# Iterate over metadata_ attributes and fetch their values
			metadata = {} # metadata_ -> dict
			for item in metadata_:
				k_ = item.key.lower()
				for s in ('time', 'date', 'modified', 'created'):
					if s in k_ and item.values:
						metadata[item.key] = item.values[0].value
						break # for s in >>>>>>>>>>>>>>>>>
			#
			#offs_ = metadata.get('offset_time_original','') # this is missing in Hachoir?; was introduced in Exif v.2.31(2016)
			#eq. OffsetTimeOriginal: '+02:00' 
			for (k0, k1) in {'creation_date': 'created', 'date_time_digitized': 'modified'}.items():
				v_ = metadata.get(k0)
				if v_:
					metadata[k1] = v_
			return add_missing_modified(metadata)

		def add_missing_modified(metadata):
			if add_modified_if_missing: # mainly 'modified' will be used to correct the timestamp of the file
				if not metadata.get('modified', None):
					v_ = metadata.get('created', None)
					if v_:
						metadata['modified'] = v_ 
						if FromMetadata.VERBOSE: print('    c --> m')
			return metadata

		### 
		f_path = f'{self.file_path}'
		path_end = f_path[-8:].lower()
		try:
			if path_end.endswith('.pdf'):
				return extract_pdf_metadata(f_path)
			elif path_end.endswith(('.docx', '.xlsx', '.pptx')):
				return extract_office_metadata(f_path)
			else: # all other: hachoir
				return extract_multimedia_metadata(f_path)
		except Exception as error:
			if FromMetadata.VERBOSE: print(f'{error}')
			return {}

#-=-=-=-#

	def set_date(self, cm_='cm'):
		"""
		Sets the date based on file metadata
		"""
		self._metadata = self.extract_metadata()	#$# ;print(f'\t#$# {self._metadata=}') #$#
		if not self._metadata:
			return None
		
		if FromMetadata.VERBOSE:
			for k,v in self._metadata.items():
				k_ = k.lower()
				for s in ('time', 'date', 'modified', 'created'):
					if s in k_:
						if k_ in ('modified', 'created'):
							if isinstance(v, datetime) and (v.year >= 1970): # .astimezone() problem with 1601-01-01 00:00+02:00
								print(f'    {k}: {v} ({v.astimezone()}, tzinfo: {v.tzinfo})')
							elif isinstance(v, str):
								print(f'    {k}: {v} ({parse(v).astimezone()})')
						else:
							print(f'    {k}: {v}')
						break # for s in >>>>>>>>>
			
		dt_tm_param = {'created': None, 'modified': None}
		for k in dt_tm_param.keys():
			if k[0] in cm_:
				dt_tm_param[k] = self._metadata.get(k)
		
		# return str(FileDate.self) if not FileDate.SET_SILENT
		return FileDate(self.file_path).set(**dt_tm_param) if not FromMetadata.LIST_ONLY else None

#-=-=-=-#

def more_tests():
	"""
	My test files in development...
	>>> FromMetadata.VERBOSE = 1; FromMetadata.LIST_ONLY = 0
	>>> for file_path in ("ab c.pdf", "ab cc.pdf", "ab ccc.pdf", "ab c.docx", "ab c.xlsx", "ab c.pptx", "ab c.jpg", "ab c.exe", "ab cc.exe" ):
	...   fp = "tmp/"+	file_path
	...   ''; f'''{fp} ================='''
	...   FileDate.SET_SILENT = 1; FileDate(fp).set(created= '2000-01-02 12:00:00', modified= '2021-03-04 14:00:00')
	...   FileDate.SET_SILENT = 0; f'''{FromMetadata(fp).set_date('cm')}'''
	''
	'tmp/ab c.pdf ================='
	    CreationDate: D:20240528193019+00'00'
	    ModDate: D:20240528193019+00'00'
	    created: 2024-05-28 19:30:19+00:00 (2024-05-28 21:30:19+02:00, tzinfo: tzutc())
	    modified: 2024-05-28 19:30:19+00:00 (2024-05-28 21:30:19+02:00, tzinfo: tzutc())
	"{'created': '2024-05-28 21:30:19', 'modified': '2024-05-28 21:30:19', 'accessed': '...'}"
	''
	'tmp/ab cc.pdf ================='
	    c --> m
	    CreationDate: D:20230914214404Z
	    created: 2023-09-14 21:44:04+00:00 (2023-09-14 23:44:04+02:00, tzinfo: tzutc())
	    modified: 2023-09-14 21:44:04+00:00 (2023-09-14 23:44:04+02:00, tzinfo: tzutc())
	"{'created': '2023-09-14 23:44:04', 'modified': '2023-09-14 23:44:04', 'accessed': '...'}"
	''
	'tmp/ab ccc.pdf ================='
	    c --> m
	    CreationDate: D:20210316065455Z00'00'
	    created: 2021-03-16 06:54:55+00:00 (2021-03-16 07:54:55+01:00, tzinfo: tzutc())
	    modified: 2021-03-16 06:54:55+00:00 (2021-03-16 07:54:55+01:00, tzinfo: tzutc())
	"{'created': '2021-03-16 07:54:55', 'modified': '2021-03-16 07:54:55', 'accessed': '...'}"
	''
	'tmp/ab c.docx ================='
	    created: 2021-04-07 06:21:00+00:00 (2021-04-07 08:21:00+02:00, tzinfo: UTC)
	    modified: 2023-10-03 05:14:00+00:00 (2023-10-03 07:14:00+02:00, tzinfo: UTC)
	"{'created': '2021-04-07 08:21:00', 'modified': '2023-10-03 07:14:00', 'accessed': '...'}"
	''
	'tmp/ab c.xlsx ================='
	    ### missing tzinfo +00 ###
	    ### missing tzinfo +00 ###
	    created: 2009-10-27 10:23:20+00:00 (2009-10-27 11:23:20+01:00)
	    modified: 2023-02-02 10:00:06+00:00 (2023-02-02 11:00:06+01:00)
	"{'created': '2009-10-27 11:23:20', 'modified': '2023-02-02 11:00:06', 'accessed': '...'}"
	''
	'tmp/ab c.pptx ================='
	    ### missing tzinfo +00 ###
	    ### missing tzinfo +00 ###
	    created: 2013-01-27 09:14:16+00:00 (2013-01-27 10:14:16+01:00)
	    modified: 2013-01-27 09:15:58+00:00 (2013-01-27 10:15:58+01:00)
	"{'created': '2013-01-27 10:14:16', 'modified': '2013-01-27 10:15:58', 'accessed': '...'}"
	''
	'tmp/ab c.jpg ================='
	    creation_date: 2024-05-21 13:49:27
	    date_time_original: 2024-05-21 13:49:27
	    date_time_digitized: 2024-05-21 13:49:27
	    created: 2024-05-21 13:49:27 (2024-05-21 13:49:27+02:00, tzinfo: None)
	    modified: 2024-05-21 13:49:27 (2024-05-21 13:49:27+02:00, tzinfo: None)
	"{'created': '2024-05-21 13:49:27', 'modified': '2024-05-21 13:49:27', 'accessed': '...'}"
	''
	'tmp/ab c.exe ================='
	    c --> m
	    creation_date: 2017-05-03 08:30:12
	    created: 2017-05-03 08:30:12 (2017-05-03 08:30:12+02:00, tzinfo: None)
	    modified: 2017-05-03 08:30:12 (2017-05-03 08:30:12+02:00, tzinfo: None)
	"{'created': '2017-05-03 08:30:12', 'modified': '2017-05-03 08:30:12', 'accessed': '...'}"
	''
	'tmp/ab cc.exe ================='
	'None'
"""


#=========================================

if __name__ == '__main__':
	import doctest
	#doctest.testmod(verbose=True, optionflags=doctest.ELLIPSIS)
	doctest.testmod(optionflags=doctest.ELLIPSIS)
